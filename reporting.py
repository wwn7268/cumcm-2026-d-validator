"""Local report export; student input files are never modified."""
from pathlib import Path
from datetime import datetime
import csv,json,re
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from visuals import make_figure

STATUS={'feasible':'可行','infeasible':'不可行','invalid':'输入错误或校验未完成'}
METRIC_LABELS={'executed_count':'执行装备数','rectangle_count':'使用矩形数','added_count':'新增C类数量',
              'occupied_area':'实际占用并集面积','demand_area':'各次占用面积之和','conflicting_cells':'冲突单元数',
              'conflict_pairs':'冲突装备对数','conflict_events':'交叠事件数','max_overlap':'最大同时占用数',
              'base_occupied_area':'第二问占用面积','added_occupied_area':'新增占用面积',
              'added_demand_area':'新增计划各次面积之和','utilization':'总面积利用率',
              'base_utilization':'第二问面积利用率','utilization_increase':'新增面积占总资源比例',
              'remaining_resource_utilization':'新增面积占原剩余资源比例','window_area':'资源窗口面积',
              'objective_tuple':'原有装备九项指标','objective_six':'原有装备前六项指标',
              'frequency_adjusted':'调频装备数','time_adjusted':'调时装备数','gap_adjusted':'改间隔装备数',
              'by_category':'原有装备分类统计','analysis_complete':'几何检查完成',
              'overlap_excess_area':'重复占用面积','conflicts_truncated':'冲突明细是否截断',
              'conflict_counts':'冲突类型统计','max_end':'最后结束时刻','min_start':'最早开始时刻',
              'metrics_valid_for_feasible_plan':'指标是否来自可行方案'}

def text_of(item):
    if isinstance(item,dict):return str(item.get('message',item))
    return str(item)

def summary_text(result):
    settings=result.get('settings',{});s=result.get('summary',{})
    lines=[f'问题{result.get("question")}：{STATUS.get(result.get("status"),result.get("status","未知"))}',
           f'验证时间窗口：[0,{settings.get("horizon",643)})；频率窗口：[0,100)',
           '判定范围：附件格式、参数规则、完整重复使用计划和时频冲突；不判定最优性。']
    if result.get('question')==3:lines.append('第三问固定所选第二问执行方案，已有装备的位置不再改变。')
    if result.get('question')==4:
        lines+=['第四问以原始150台装备为基准；仅C类可改间隔，间隔变化量不超过10Δt且新间隔非负。',
                '调频、调时、改间隔至多选择一项；改间隔时首次开始保持不变，后续使用按新周期展开。']
    if settings.get('input_file'):lines.append('提交附件：'+settings['input_file'])
    if settings.get('base_file'):lines.append('第二问基础附件：'+settings['base_file'])
    lines+=['','统计：']
    labels={k:METRIC_LABELS[k] for k in ('executed_count','rectangle_count','added_count','occupied_area','demand_area',
            'conflicting_cells','conflict_pairs','conflict_events','max_overlap','utilization',
            'base_occupied_area','added_occupied_area','base_utilization','utilization_increase','remaining_resource_utilization',
            'frequency_adjusted','time_adjusted','gap_adjusted')}
    for key,label in labels.items():
        if key in s:
            value=s[key]
            if ('utilization' in key) and isinstance(value,(int,float)):value=f'{100*value:.2f}%'
            lines.append(f'  {label}：{value}')
    objective=s.get('objective_tuple',result.get('objective_tuple'))
    if objective:
        title={2:'第二问指标',3:'第二问基础方案指标',4:'第四问指标'}.get(result.get('question'),'原有装备指标')
        lines+=['',title+'（幅度按基本单位等权）：',
                '  顺序：总撤销、A撤销、B撤销、总调整、A调整、B调整、总幅度、A幅度、B幅度',
                '  '+str(tuple(objective)), '  指标用于比较方案，不构成最优性证明。']
        if result.get('question')==4:
            lines.append('  仅改间隔的装备计1次调整；幅度计新旧间隔之差的绝对值，不累计各次使用的位移。')
    counts=s.get('by_category',s.get('counts'))
    if counts:
        lines+=['','分类统计：']
        for name,values in counts.items():
            if isinstance(values,dict):
                labels={'total':'原有总数','kept':'原样','adjusted':'调整','canceled':'撤销','amplitude':'调整幅度',
                        'frequency_adjusted':'调频','time_adjusted':'调时','gap_adjusted':'改间隔'}
                lines.append(f'  {name}类：'+ '，'.join(f'{label}{values[key]}' for key,label in labels.items() if key in values))
            else:lines.append(f'  {name}：{values}')
    errors=result.get('errors',[]);conflicts=result.get('conflicts',[])
    lines+=['',f'规则或格式问题：{len(errors)}条；展示的交叠事件：{len(conflicts)}条。']
    for error in errors[:25]:
        if isinstance(error,dict):
            loc=' / '.join(str(error.get(k,'')) for k in ('sheet','row','id') if error.get(k) not in ('',None))
            lines.append(f'  {loc+"：" if loc else ""}{text_of(error)}')
        else:lines.append('  '+str(error))
    if len(errors)>25:lines.append('  更多问题见“问题与冲突”页和导出报告。')
    if result.get('warnings'):
        lines+=['','说明：']+[f'  {text_of(w)}' for w in result['warnings']]
    lines+=['','面积利用率以实际占用区域的并集面积为分子，空闲间隔不计入。',
            '默认时间上限643沿用原计划最晚结束时刻；这是验证设置，应与论文的假设一致。']
    return '\n'.join(lines)

def safe_value(v):
    if isinstance(v,(dict,list,tuple)):v=json.dumps(v,ensure_ascii=False)
    if isinstance(v,str) and v[:1] in ('=','+','-','@','\t','\r'):return "'"+v
    return v

def export_report(result, directory):
    directory=Path(directory);directory.mkdir(parents=True,exist_ok=True)
    # Use distinct folders at the UI/CLI entry; never overwrite the student's workbook.
    (directory/'验证报告.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    (directory/'验证摘要.txt').write_text(summary_text(result),encoding='utf-8-sig')
    figure,counts,categories=make_figure(result)
    figure.savefig(directory/'时频占用矩阵.png',dpi=170)
    figure.clear()
    with (directory/'占用计数矩阵.csv').open('w',encoding='utf-8-sig',newline='') as f:
        writer=csv.writer(f);writer.writerow(['时间单元起点']+[f'频段[{j},{j+1})' for j in range(100)])
        for t,row in enumerate(counts):writer.writerow([t,*map(int,row)])
    wb=Workbook();overview=wb.active;overview.title='验证摘要'
    overview.append(['项目','结果'])
    overview.append(['判定',STATUS.get(result.get('status'),result.get('status'))])
    overview.append(['问题',result.get('question')])
    overview.append(['验证范围','可行性；不证明最优性'])
    for key,value in result.get('summary',{}).items():
        if 'utilization' in key and isinstance(value,(int,float)):value=f'{value:.6%}'
        overview.append([METRIC_LABELS.get(key,key),safe_value(value)])
    for key,value in result.get('settings',{}).items():overview.append([key,safe_value(value)])
    error_ws=wb.create_sheet('格式与规则问题');error_ws.append(['代码','工作表','行号','装备或序号','问题'])
    for error in result.get('errors',[]):
        if isinstance(error,dict):error_ws.append([safe_value(error.get(k,'')) for k in ('code','sheet','row','id','message')])
        else:error_ws.append(['','','','',safe_value(str(error))])
    conflict_ws=wb.create_sheet('时频冲突');conflict_ws.append(['类型','装备1','第几次1','装备2','第几次2','频段下界','频段上界','时间下界','时间上界'])
    for c in result.get('conflicts',[]):conflict_ws.append([safe_value(c.get(k,'')) for k in ('kind','id1','use1','id2','use2','f0','f1','t0','t1')])
    plan_ws=wb.create_sheet('重建执行计划');keys=['id','cls','f0','f1','t0','t1','gap','count','canceled','action','is_new','df','dt','dg']
    plan_ws.append(['编号','类别','频段下界','频段上界','首次开始','首次结束','空闲间隔','使用次数','撤销','动作','新增','频率变化量','首次时间变化量','间隔变化量'])
    for p in result.get('plans',[]):
        changes='+'.join(label for key,label in (('df','调频'),('dt','调时'),('dg','改间隔')) if p.get(key))
        action='撤销' if p.get('canceled') else '新增' if p.get('is_new') else changes or '原样'
        plan_ws.append([safe_value(action if k=='action' else p.get(k,'')) for k in keys])
    for ws in wb:
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]:cell.font=Font(name='Microsoft YaHei',bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='245F87')
        for col in ws.columns:
            letter=col[0].column_letter
            longest=max((len(str(c.value or '')) for c in col),default=10)
            ws.column_dimensions[letter].width=min(80,max(14,longest*1.2))
        for row in ws:
            for c in row:c.alignment=Alignment(vertical='top',wrap_text=True)
    wb.save(directory/'验证明细.xlsx')
    return directory

def new_output_dir(root,submission,question):
    safe=re.sub(r'[^\w\-\u4e00-\u9fff]+','_',Path(submission).stem)[:70]
    return Path(root)/f'{datetime.now():%Y%m%d_%H%M%S_%f}_问题{question}_{safe}'
