"""D 题问题二、三、四附件的独立可行性检查。原始附件只读，不运行优化求解器。"""
from __future__ import annotations

import csv
import io
import json
import math
import re
import unicodedata
from pathlib import Path
from typing import Any

import numpy as np
from openpyxl import load_workbook

DEFAULT_ORIGINAL = Path(__file__).resolve().parent / "data" / "original_plans.json"
MAX_BYTES = 25 * 1024 * 1024
MAX_ROWS = 2000
MAX_COLUMNS = 40
MAX_RECTANGLES = 25000
MAX_CONFLICTS = 50000


class InputError(ValueError):
    def __init__(self, code: str, message: str, **location: Any):
        super().__init__(message)
        self.code = code
        self.location = location


def _error(code: str, message: str, *, sheet=None, row=None, id=None) -> dict:
    return dict(code=code, message=message, sheet=sheet, row=row, id=id)


def _text(value: Any) -> str:
    return unicodedata.normalize("NFKC", str(value)).strip()


def _blank(value: Any) -> bool:
    return value is None or isinstance(value, str) and not value.strip()


def _integer(value: Any, label: str) -> int:
    if isinstance(value, bool) or _blank(value):
        raise InputError("invalid_integer", f"{label}必须是整数。")
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        if math.isfinite(value) and value.is_integer():
            return int(value)
        raise InputError("invalid_integer", f"{label}必须是有限整数。")
    text = _text(value)
    if not re.fullmatch(r"[+-]?\d+(?:\.0+)?", text):
        raise InputError("invalid_integer", f"{label}必须是整数，收到：{text[:80]}。")
    return int(text.split(".")[0])


def _interval(value: Any, label: str) -> tuple[int, int]:
    if _blank(value):
        raise InputError("missing_interval", f"缺少{label}。")
    text = _text(value).replace("，", ",").replace("−", "-")
    # 区间端点必须显式给出；容许常见的闭括号写法，统一按半开区间解释。
    match = re.fullmatch(r"[\[(]\s*([+-]?\d+(?:\.0+)?)\s*[,;]\s*([+-]?\d+(?:\.0+)?)\s*[\])]", text)
    if not match:
        raise InputError("invalid_interval", f"{label}格式应为 [起点,终点)，例如 [10,20)；收到：{text[:80]}。")
    left = _integer(match.group(1), label + "起点")
    right = _integer(match.group(2), label + "终点")
    if right <= left:
        raise InputError("invalid_interval", f"{label}终点必须大于起点。")
    return left, right


def _cancel(value: Any) -> bool:
    if _blank(value):
        return False
    if isinstance(value, bool):
        return value
    text = _text(value).lower() if not _blank(value) else ""
    if text in {"是", "撤销", "已撤销", "取消", "1", "1.0", "true", "yes", "y"}:
        return True
    if text in {"否", "不撤销", "保留", "执行", "0", "0.0", "false", "no", "n"}:
        return False
    raise InputError("invalid_cancel_flag", "是否撤销用频计划必须明确填写“是/否”（或 1/0）。")


_ALIASES = {
    "id": {"用频装备编号", "装备编号", "设备编号", "id"},
    "serial": {"新增用频装备序号", "新增装备序号", "新增设备序号", "序号"},
    "frequency": {"调整后频段区间", "频段区间", "调整后频率区间", "频率区间", "调整后频段范围", "频段范围"},
    "time": {"调整后时间区间", "时间区间", "首次时间区间", "首次用频时间区间"},
    "cancel": {"是否撤销用频计划", "是否撤销", "是否取消"},
    "gap": {"gap", "时间间隔", "用频时间间隔", "用频间隔", "间隔", "间隔时间", "间隔时长", "调整后时间间隔", "调整后用频间隔", "调整后间隔时长", "调整后间隔时间"},
    "count": {"count", "使用次数", "用频次数", "重复次数", "次数", "调整后使用次数", "调整后用频次数"},
    "width": {"带宽", "频段宽度", "用频带宽", "调整后带宽", "width"},
    "duration": {"使用时长", "用频时长", "每次用频时长", "持续时间", "持续时长", "duration"},
    "cls": {"类别", "装备类别", "设备类别", "用频装备类别", "类型", "cls"},
    "note": {"备注", "说明", "note"},
}


def _header_key(value: Any) -> str:
    return re.sub(r"\s+", "", _text(value)).lower()


_HEADER_MAP = {_header_key(alias): name for name, aliases in _ALIASES.items() for alias in aliases}


def _check_file(path: Path) -> None:
    if not path.is_file():
        raise InputError("file_not_found", f"附件不存在：{path}")
    if path.stat().st_size > MAX_BYTES:
        raise InputError("file_limit", f"附件超过 {MAX_BYTES // 1024 // 1024} MB 的检查上限，未完成检查。")
    if path.suffix.lower() not in {".xlsx", ".xlsm", ".csv"}:
        raise InputError("unsupported_file", "支持 .xlsx、.xlsm 和 UTF-8/GB18030 编码的 .csv；旧 .xls 请另存为 .xlsx。")


def _csv_rows(path: Path) -> list[list[Any]]:
    raw = path.read_bytes()
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            text = raw.decode("gb18030")
        except UnicodeDecodeError as exc:
            raise InputError("csv_encoding", "CSV 编码无法识别，请另存为 UTF-8 CSV。") from exc
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    rows = []
    for row in csv.reader(io.StringIO(text), dialect):
        if len(rows) > MAX_ROWS + 20 or len(row) > MAX_COLUMNS:
            raise InputError("table_limit", f"附件最多允许 {MAX_ROWS} 条数据、{MAX_COLUMNS} 列，未完成检查。")
        rows.append(row)
    return rows


def read_sheet_names(path: Path | str) -> list[str]:
    path = Path(path)
    _check_file(path)
    if path.suffix.lower() == ".csv":
        return ["CSV"]
    workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


def _header(rows: list[list[Any]], question: int) -> tuple[int, dict[str, int]] | None:
    required = {"id", "frequency", "time", "cancel"} if question in (2, 4) else {"serial", "frequency", "time"}
    if question == 4:
        required.add("gap")
    for number, row in enumerate(rows[:20]):
        mapping = {}
        for col, value in enumerate(row):
            if not _blank(value) and (key := _HEADER_MAP.get(_header_key(value))):
                if key in mapping and key != "note":
                    mapping["__duplicate__"] = col
                mapping[key] = col
        if required <= mapping.keys():
            return number, mapping
    return None


def _read_table(path: Path, question: int, sheet: str | None) -> tuple[str, list[list[Any]], int, dict[str, int]]:
    _check_file(path)
    if path.suffix.lower() == ".csv":
        if sheet not in (None, "", "CSV"):
            raise InputError("sheet_not_found", "CSV 的表名为 CSV。")
        rows = _csv_rows(path)
        found = _header(rows, question)
        if not found:
            raise InputError("missing_header", f"前 20 行未找到问题{question}的完整附件表头。", sheet="CSV")
        return "CSV", rows, *found
    try:
        workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    except Exception as exc:
        raise InputError("unreadable_workbook", f"无法打开 Excel：{exc}") from exc
    try:
        if sheet and sheet not in workbook.sheetnames:
            raise InputError("sheet_not_found", f"找不到工作表：{sheet}")
        names = [sheet] if sheet else workbook.sheetnames
        matched = []
        for name in names:
            ws = workbook[name]
            # 只读取少量表头，选定后才读取数据；不执行任何公式或宏。
            preview = [[c.value for c in row] for row in ws.iter_rows(min_row=1, max_row=20, max_col=min(ws.max_column or 1, MAX_COLUMNS + 1))]
            found = _header(preview, question)
            if found:
                matched.append((name, found))
        if not matched:
            raise InputError("missing_header", f"未找到问题{question}的完整附件表头。", sheet=sheet)
        if len(matched) > 1:
            raise InputError("ambiguous_sheet", "多个工作表符合附件格式，请明确选择其中一个：" + "、".join(x[0] for x in matched))
        name, found = matched[0]
        ws = workbook[name]
        if (ws.max_row or 0) > MAX_ROWS + 20 or (ws.max_column or 0) > MAX_COLUMNS:
            raise InputError("table_limit", f"工作表范围超过 {MAX_ROWS} 条数据或 {MAX_COLUMNS} 列，请去掉多余格式行后重试；未完成检查。", sheet=name)
        rows = []
        for rowno, cells in enumerate(ws.iter_rows(), 1):
            values = []
            for cell in cells:
                if cell.data_type == "f":
                    raise InputError("formula_cell", f"单元格 {cell.coordinate} 含公式；请复制后粘贴为数值再提交。", sheet=name, row=rowno)
                if cell.data_type == "e":
                    raise InputError("excel_error_cell", f"单元格 {cell.coordinate} 是 Excel 错误值。", sheet=name, row=rowno)
                values.append(cell.value)
            rows.append(values)
        return name, rows, *found
    finally:
        workbook.close()


def _load_original(path: Path) -> dict[str, dict]:
    try:
        data = json.loads(path.read_text(encoding="utf-8-sig"))
        out = {}
        for record in data:
            plan = {key: record[key] for key in ("id", "cls", "f0", "f1", "t0", "t1", "gap", "count")}
            for key in ("f0", "f1", "t0", "t1", "gap", "count"):
                plan[key] = _integer(plan[key], "原始数据 " + key)
            if plan["id"] in out or plan["cls"] not in {"A", "B", "C"}:
                raise ValueError("编号或类别错误")
            plan.update(canceled=False, is_new=False, source_sheet=None, source_row=None, df=0, dt=0, dg=0)
            out[plan["id"]] = plan
        if len(out) != 150 or {c: sum(p["cls"] == c for p in out.values()) for c in "ABC"} != {"A": 20, "B": 40, "C": 90}:
            raise ValueError("原始装备数量应为 A20、B40、C90")
        return out
    except Exception as exc:
        raise InputError("original_data_error", f"原始数据无法读取或内容不完整：{exc}") from exc


def _normalize_id(value: Any, question: int) -> str:
    text = _text(value).upper().replace(" ", "")
    if question in (2, 4):
        found = re.fullmatch(r"([ABC])0*(\d+)", text)
        if not found:
            raise InputError("invalid_id", f"装备编号应为 A001、B001、C001 等格式，收到：{text[:80]}。")
        return f"{found.group(1)}{int(found.group(2)):03d}"
    found = re.fullmatch(r"(?:NEWC|C)?0*(\d+)(?:\.0+)?", text)
    if not found or int(found.group(1)) < 1:
        raise InputError("invalid_serial", "新增装备序号必须为正整数，也可填写 C001、NEWC001 等编号。")
    return f"NEWC{int(found.group(1)):03d}"


def _parse_rows(path: Path, question: int, original: dict[str, dict], sheet: str | None, horizon: int) -> tuple[list[dict], list[dict], list[str], bool]:
    name, rows, header_no, columns = _read_table(path, question, sheet)
    errors, warnings = [], []
    incomplete = False
    if "__duplicate__" in columns:
        raise InputError("duplicate_header", "表头包含重复字段，请保留一列。", sheet=name, row=header_no + 1)
    if len(rows) - header_no - 1 > MAX_ROWS:
        raise InputError("table_limit", f"最多允许 {MAX_ROWS} 条数据；未完成检查。", sheet=name)
    known_columns = set(columns.values())
    for col, header in enumerate(rows[header_no]):
        if col not in known_columns and not _blank(header):
            errors.append(_error("unsupported_column", f"无法解释字段“{_text(header)}”；请使用官方表头，不能将未知字段忽略后判为通过。", sheet=name, row=header_no + 1))
            incomplete = True
    if any(any(not _blank(v) for v in row) for row in rows[:header_no]):
        warnings.append(f"工作表 {name} 的表头位于第 {header_no + 1} 行；其上方文字不属于计划数据。")
    plans = {key: dict(value) for key, value in original.items()} if question in (2, 4) else {}
    seen = set()
    used_implicit_unchanged = False
    for rowno, row in enumerate(rows[header_no + 1:], header_no + 2):
        if all(_blank(v) for v in row):
            continue
        row_errors = []
        pid = None
        def fail(code: str, message: str, invalid=False):
            nonlocal incomplete
            row_errors.append(_error(code, message, sheet=name, row=rowno, id=pid))
            incomplete = incomplete or invalid
        def value(key: str):
            col = columns.get(key)
            return row[col] if col is not None and col < len(row) else None
        if any(isinstance(v, str) and v.lstrip().startswith("=") for v in row):
            fail("formula_cell", "数据行含公式，不能使用缓存值代替检查；请粘贴为数值。", True)
            errors.extend(row_errors)
            continue
        if any(not _blank(v) for col, v in enumerate(row) if col not in known_columns):
            fail("unlabeled_data", "数据出现在没有可识别表头的列中。", True)
        try:
            pid = _normalize_id(value("id" if question in (2, 4) else "serial"), question)
        except InputError as exc:
            fail(exc.code, str(exc), True)
            errors.extend(row_errors)
            continue
        if pid in seen:
            fail("duplicate_id", f"编号 {pid} 重复出现，无法确定应执行哪一条计划。", True)
            errors.extend(row_errors)
            continue
        seen.add(pid)
        if question in (2, 4) and pid not in original:
            fail("unknown_id", f"原始附件中不存在装备 {pid}。", True)
            errors.extend(row_errors)
            continue
        origin = original[pid] if question in (2, 4) else dict(id=pid, cls="C", gap=8, count=12)
        plan = dict(origin)
        plan.update(canceled=False, is_new=question == 3, source_sheet=name, source_row=rowno, df=0, dt=0, dg=0)
        try:
            plan["canceled"] = _cancel(value("cancel")) if question in (2, 4) else False
            if question in (2, 4) and _blank(value("cancel")):
                used_implicit_unchanged = True
            if question == 3 and "cancel" in columns and not _blank(value("cancel")) and _cancel(value("cancel")):
                fail("q3_cancellation", "问题三的附件应列新增计划，不能使用撤销行抵扣已有计划。")
            for key in ("frequency", "time"):
                if question in (2, 4) and _blank(value(key)):
                    used_implicit_unchanged = True
                    continue
                start, end = _interval(value(key), "频段区间" if key == "frequency" else "时间区间")
                plan["f0" if key == "frequency" else "t0"] = start
                plan["f1" if key == "frequency" else "t1"] = end
            for key in ("gap", "count"):
                if key in columns and not _blank(value(key)):
                    candidate = _integer(value(key), "用频间隔" if key == "gap" else "使用次数")
                    if key == "gap" and question == 4:
                        plan["dg"] = candidate - origin[key]
                        if plan["cls"] != "C" and plan["dg"]:
                            fail("gap_class_limit", f"{pid} 属于 {plan['cls']} 类，第四问仅允许 C 类调整用频间隔。")
                        if abs(plan["dg"]) > 10:
                            fail("gap_shift_limit", f"{pid} 用频间隔变化 {plan['dg']}，超过 ±10。")
                    elif candidate != origin[key]:
                        fail("changed_" + key, f"{pid} 的{('用频间隔' if key == 'gap' else '使用次数')}必须保持为 {origin[key]}，收到 {candidate}。")
                    plan[key] = candidate
            width = plan["f1"] - plan["f0"]
            duration = plan["t1"] - plan["t0"]
            expected_width = origin["f1"] - origin["f0"] if question in (2, 4) else 3
            expected_duration = origin["t1"] - origin["t0"] if question in (2, 4) else 2
            if width != expected_width:
                fail("changed_width", f"{pid} 带宽应为 {expected_width}，收到 {width}。")
            if duration != expected_duration:
                fail("changed_duration", f"{pid} 每次时长应为 {expected_duration}，收到 {duration}。")
            for key, expected in (("width", expected_width), ("duration", expected_duration)):
                if key in columns and not _blank(value(key)) and _integer(value(key), key) != expected:
                    fail("inconsistent_" + key, f"额外字段 {key} 与该类装备规定值 {expected} 不一致。")
            if "cls" in columns and not _blank(value("cls")) and _text(value("cls")).upper().replace("类", "") != plan["cls"]:
                fail("changed_class", f"{pid} 的类别应为 {plan['cls']}。")
            if question in (2, 4):
                plan["df"] = plan["f0"] - origin["f0"]
                plan["dt"] = plan["t0"] - origin["t0"]
                changed = int(plan["df"] != 0) + int(plan["dt"] != 0) + int(plan["dg"] != 0)
                if not plan["canceled"] and changed > 1:
                    names = "频段、首次时间或用频间隔" if question == 4 else "频段和首次时间"
                    fail("multiple_actions", f"{pid} 同时调整{names}中的多个参数，不符合每台最多调整一个参数的规定。")
                if abs(plan["df"]) > 10:
                    fail("frequency_shift_limit", f"{pid} 频段平移 {plan['df']}，超过 ±10。")
                if abs(plan["dt"]) > 5:
                    fail("time_shift_limit", f"{pid} 首次时间平移 {plan['dt']}，超过 ±5。")
                if plan["canceled"] and (changed or plan["gap"] != origin["gap"] or plan["count"] != origin["count"]):
                    fail("canceled_and_adjusted", f"{pid} 同时撤销和调整；撤销行的区间应留空或填写原区间。")
            if plan["gap"] < 0 or not 1 <= plan["count"] <= 1000:
                fail("invalid_repetition", "用频间隔须非负，次数须在 1..1000 的检查范围内。", True)
                # 此类数据没有可完整展开的周期，不交给图形层。
                plan = None
            elif not plan["canceled"]:
                last = plan["t0"] + (plan["count"] - 1) * (duration + plan["gap"]) + duration
                if plan["f0"] < 0 or plan["f1"] > 100:
                    fail("frequency_boundary", f"{pid} 频段 [{plan['f0']},{plan['f1']}) 超出 [0,100)。")
                if plan["t0"] < 0 or last > horizon:
                    fail("time_boundary", f"{pid} 完整周期范围 [{plan['t0']},{last}) 超出所选时间窗 [0,{horizon})。")
            if plan is not None:
                plans[pid] = plan
        except InputError as exc:
            fail(exc.code, str(exc), True)
        errors.extend(row_errors)
    if used_implicit_unchanged:
        label = "问题四的空白区间和间隔按原值保留" if question == 4 else "问题二的空白区间按原区间保留"
        warnings.append(label + "；空白撤销标志按不撤销处理，明确写“是”的行按撤销处理。")
    return list(plans.values()), errors, warnings, incomplete


def analyze_plans(plans: list[dict], horizon: int = 643) -> dict:
    """展开每次占用，计算窗口内并集面积及真实周期冲突，不裁剪冲突判定。"""
    if isinstance(horizon, bool) or not isinstance(horizon, int) or not 1 <= horizon <= 10000:
        raise ValueError("horizon 必须是 1..10000 的整数")
    errors, rectangles = [], []
    demand_area = 0
    for plan_no, p in enumerate(plans):
        if p.get("canceled", False):
            continue
        width = p["f1"] - p["f0"]
        duration = p["t1"] - p["t0"]
        count, gap = p["count"], p["gap"]
        if width <= 0 or duration <= 0 or gap < 0 or not 1 <= count <= 1000:
            errors.append(_error("invalid_geometry", "计划的几何参数无法完整展开。", sheet=p.get("source_sheet"), row=p.get("source_row"), id=p["id"]))
            continue
        # 全部执行计划都须满足所选窗口，包括附件中未列出、从原计划继承的装备。
        last = p["t0"] + (count - 1) * (duration + gap) + duration
        if p["f0"] < 0 or p["f1"] > 100:
            errors.append(_error("frequency_boundary", f"{p['id']} 频段 [{p['f0']},{p['f1']}) 超出 [0,100)。", sheet=p.get("source_sheet"), row=p.get("source_row"), id=p["id"]))
        if p["t0"] < 0 or last > horizon:
            errors.append(_error("time_boundary", f"{p['id']} 完整周期范围 [{p['t0']},{last}) 超出所选时间窗 [0,{horizon})。", sheet=p.get("source_sheet"), row=p.get("source_row"), id=p["id"]))
        if len(rectangles) + count > MAX_RECTANGLES:
            errors.append(_error("rectangle_limit", f"周期矩形超过 {MAX_RECTANGLES} 个，未完成检查。"))
            return dict(conflicts=[], summary={"analysis_complete": False}, errors=errors)
        demand_area += width * duration * count
        for use in range(count):
            t0 = p["t0"] + use * (duration + gap)
            rectangles.append((t0, t0 + duration, p["f0"], p["f1"], plan_no, use + 1))
    difference = np.zeros((horizon + 1, 101), dtype=np.int32)
    for t0, t1, f0, f1, _, _ in rectangles:
        a, b = max(t0, 0), min(t1, horizon)
        c, d = max(f0, 0), min(f1, 100)
        if a < b and c < d:
            difference[a, c] += 1
            difference[b, c] -= 1
            difference[a, d] -= 1
            difference[b, d] += 1
    matrix = difference.cumsum(axis=0).cumsum(axis=1)[:horizon, :100]
    occupied_area = int(np.count_nonzero(matrix))
    conflicting_cells = int(np.count_nonzero(matrix > 1))
    overlap_excess = int(np.maximum(matrix - 1, 0).sum())
    conflicts, active = [], []
    truncated = False
    conflict_pairs = set()
    for rect in sorted(rectangles):
        t0, t1, f0, f1, plan_no, use = rect
        active = [other for other in active if other[1] > t0]
        for other in active:
            ot0, ot1, of0, of1, other_no, other_use = other
            if plan_no == other_no:
                continue
            left, right = max(f0, of0), min(f1, of1)
            if left >= right:
                continue
            p, q = plans[other_no], plans[plan_no]
            kind = "新增—新增" if p.get("is_new") and q.get("is_new") else "原有—新增" if p.get("is_new") or q.get("is_new") else "原有—原有"
            conflicts.append(dict(id1=p["id"], id2=q["id"], use1=other_use, use2=use, f0=left, f1=right, t0=max(t0, ot0), t1=min(t1, ot1), kind=kind))
            conflict_pairs.add(tuple(sorted((p["id"], q["id"]))))
            if len(conflicts) >= MAX_CONFLICTS:
                truncated = True
                break
        if truncated:
            errors.append(_error("conflict_limit", f"冲突事件达到 {MAX_CONFLICTS} 条；事件明细未全部展开，不能认定通过。"))
            break
        active.append(rect)
    by_category = {}
    for cls in "ABC":
        old = [p for p in plans if p.get("cls") == cls and not p.get("is_new")]
        canceled = sum(bool(p.get("canceled")) for p in old)
        adjusted = sum(not p.get("canceled") and any(p.get(key, 0) for key in ("df", "dt", "dg")) for p in old)
        amplitude = sum(abs(p.get("df", 0)) + abs(p.get("dt", 0)) + abs(p.get("dg", 0)) for p in old if not p.get("canceled"))
        by_category[cls] = dict(total=len(old), kept=len(old) - canceled - adjusted, adjusted=adjusted, canceled=canceled, amplitude=amplitude,
                                frequency_adjusted=sum(not p.get("canceled") and bool(p.get("df")) for p in old),
                                time_adjusted=sum(not p.get("canceled") and bool(p.get("dt")) for p in old),
                                gap_adjusted=sum(not p.get("canceled") and bool(p.get("dg")) for p in old))
    total_canceled = sum(x["canceled"] for x in by_category.values())
    total_adjusted = sum(x["adjusted"] for x in by_category.values())
    total_amplitude = sum(x["amplitude"] for x in by_category.values())
    a, b = by_category["A"], by_category["B"]
    objective = [total_canceled, a["canceled"], b["canceled"], total_adjusted, a["adjusted"], b["adjusted"], total_amplitude, a["amplitude"], b["amplitude"]]
    return dict(conflicts=conflicts, errors=errors, summary=dict(
        analysis_complete=not any(e["code"] in {"invalid_geometry", "rectangle_limit", "conflict_limit"} for e in errors), objective_tuple=objective, objective_six=objective[:6], by_category=by_category,
        added_count=sum(bool(p.get("is_new")) and not p.get("canceled") for p in plans),
        frequency_adjusted=sum(x["frequency_adjusted"] for x in by_category.values()),
        time_adjusted=sum(x["time_adjusted"] for x in by_category.values()),
        gap_adjusted=sum(x["gap_adjusted"] for x in by_category.values()),
        executed_count=sum(not p.get("canceled") for p in plans), rectangle_count=len(rectangles),
        demand_area=int(demand_area), occupied_area=occupied_area, conflicting_cells=conflicting_cells,
        overlap_excess_area=overlap_excess, max_overlap=int(matrix.max(initial=0)),
        window_area=100 * horizon, utilization=occupied_area / (100 * horizon),
        conflict_events=len(conflicts), conflict_pairs=len(conflict_pairs), conflicts_truncated=truncated,
        conflict_counts={kind: sum(c["kind"] == kind for c in conflicts) for kind in ("原有—原有", "原有—新增", "新增—新增")},
        max_end=max((r[1] for r in rectangles), default=0), min_start=min((r[0] for r in rectangles), default=0),
    ))


def validate_submission(question: int, submission: Path | str, base_q2: Path | str | None = None, *,
                        horizon: int = 643, sheet: str | None = None, base_sheet: str | None = None,
                        original_path: Path | str | None = None) -> dict:
    """验证提交附件。问题三需要其问题二基础方案；问题四独立从原始附件重建。"""
    result = dict(question=question, status="invalid", passed=False, errors=[], warnings=[], plans=[], conflicts=[], summary={},
                  settings=dict(horizon=horizon, frequency_range=[0, 100], interval_convention="half-open", fixed_base=question == 3,
                                horizon_note="时间窗是建模约定，默认 [0,643)，可由教师修改；并非将 643 视为题目额外硬性条件。",
                                amplitude_note=("幅度按频率、首次时间与用频间隔的基本单位变化量等权相加；间隔只计一次变化量，不乘重复次数；撤销不计幅度。"
                                                if question == 4 else "幅度按频率与首次时间的基本单位变化量等权相加；撤销不计幅度。"),
                                adjustment_rule=("以原始附件 150 台装备为基准；每台最多调整频段、首次时间、用频间隔中的一项，或撤销。仅 C 类可改间隔，|Δgap|≤10 且 gap≥0；所有类别仍满足 |Δf|≤10、|Δt|≤5。"
                                                 if question == 4 else "问题二每台最多调整频段或首次时间一项，或撤销；问题三固定其对应问题二方案。"),
                                scope="只检验可行性与指标，不证明撤销/调整/幅度最小或新增数最大。",
                                input_file=str(Path(submission)), original_file=str(original_path or DEFAULT_ORIGINAL)))
    try:
        if question not in (2, 3, 4):
            raise InputError("invalid_question", "仅支持问题二、问题三或问题四。")
        if isinstance(horizon, bool) or not isinstance(horizon, int) or not 1 <= horizon <= 10000:
            raise InputError("invalid_horizon", "时间窗上限必须是 1..10000 的整数。")
        original = _load_original(Path(original_path) if original_path else DEFAULT_ORIGINAL)
        incomplete = False
        base = None
        if question == 3:
            if not base_q2:
                raise InputError("missing_base", "问题三必须提供同一套方案的问题二附件，才能检查新增计划与固定原有计划的冲突。")
            base = validate_submission(2, base_q2, horizon=horizon, sheet=base_sheet, original_path=original_path)
            result["base_validation"] = {key: base[key] for key in ("status", "passed", "errors", "summary", "settings")}
            result["settings"]["base_file"] = str(base_q2)
            if not base["passed"]:
                result["errors"].append(_error("base_not_feasible", "第二问基础方案未通过检查；问题三不能判为可行。"))
                result["errors"].extend({**e, "code": "base_" + e["code"]} for e in base["errors"])
                incomplete = base["status"] == "invalid"
        plans, errors, warnings, parse_incomplete = _parse_rows(Path(submission), question, original, sheet, horizon)
        result["errors"].extend(errors)
        result["warnings"].extend(warnings)
        incomplete = incomplete or parse_incomplete
        result["plans"] = ([dict(p) for p in base["plans"]] if base else []) + plans
        analyzed = analyze_plans(result["plans"], horizon)
        result["conflicts"] = analyzed["conflicts"]
        result["summary"] = analyzed["summary"]
        # 解析阶段已报告的同一边界问题不重复列出。
        reported = {(e["code"], e.get("id")) for e in result["errors"]}
        result["errors"].extend(e for e in analyzed["errors"] if (e["code"], e.get("id")) not in reported)
        incomplete = incomplete or not analyzed["summary"].get("analysis_complete", False)
        if result["conflicts"]:
            result["errors"].append(_error("resource_conflict", f"发现 {len(result['conflicts'])} 次时频占用冲突，详见冲突明细。"))
        if base and "occupied_area" in result["summary"] and "occupied_area" in base["summary"]:
            original_area = base["summary"]["occupied_area"]
            added_area = sum((p["f1"] - p["f0"]) * (p["t1"] - p["t0"]) * p["count"] for p in plans if not p.get("canceled"))
            result["summary"].update(base_occupied_area=original_area, added_demand_area=added_area,
                                     base_utilization=original_area / (100 * horizon))
            if not result["errors"]:
                free = 100 * horizon - original_area
                result["summary"].update(added_occupied_area=added_area,
                                         utilization_increase=added_area / (100 * horizon),
                                         remaining_resource_utilization=added_area / free if free else None)
        result["status"] = "invalid" if incomplete else "infeasible" if result["errors"] else "feasible"
        result["passed"] = result["status"] == "feasible"
        result["summary"]["metrics_valid_for_feasible_plan"] = result["passed"]
        result["warnings"].append("区间统一按左闭右开解释：端点接触不算冲突。占用面积仅统计各次实际使用，中间空档不计入。")
        if not result["passed"]:
            result["warnings"].append("图和面积统计反映可解析的数据；存在格式错误时可能不完整，存在越界时图仅显示所选窗口，不能作为有效方案成绩。")
        return result
    except InputError as exc:
        result["errors"].append(_error(exc.code, str(exc), **exc.location))
    except Exception as exc:
        # 未预料到的输入错误也必须阻止通过，交由界面展示；不执行输入内代码。
        result["errors"].append(_error("read_error", f"检查未完成：{type(exc).__name__}: {exc}"))
    return result
