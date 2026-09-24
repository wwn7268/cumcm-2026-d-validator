"""附件验证器回归测试。运行：python -m unittest discover -s tests -v"""
from pathlib import Path
import json
import sys
import unittest

from openpyxl import Workbook, load_workbook

APP = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(APP))
from engine import validate_submission  # noqa: E402

Q2_HEADER = ["用频装备编号", "调整后频段区间", "调整后时间区间", "是否撤销用频计划"]
Q3_HEADER = ["新增用频装备序号", "调整后频段区间", "调整后时间区间"]


class SubmissionTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.directory = APP / "tests" / "_generated"
        cls.directory.mkdir(parents=True, exist_ok=True)
        cls.original = json.loads((APP / "data/original_plans.json").read_text(encoding="utf-8"))
        cls.q2 = APP / "examples/result2_可行示例.xlsx"
        cls.q3 = APP / "examples/result3_可行示例.xlsx"
        cls.all_canceled_rows = [[row["id"], None, None, "是"] for row in cls.original]
        cls.all_canceled = cls.table("all_canceled.xlsx", Q2_HEADER, cls.all_canceled_rows)

    @classmethod
    def table(cls, name, header, rows):
        path = cls.directory / name
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "结果"
        sheet.append(header)
        for row in rows:
            sheet.append(row)
        workbook.save(path)
        return path

    def isolated_q2(self, replacement):
        rows = [row for row in self.all_canceled_rows if row[0] != replacement[0]]
        return self.table(self._testMethodName + ".xlsx", Q2_HEADER, rows + [replacement])

    def q3_rows(self, rows, *, base=None, **kwargs):
        path = self.table(self._testMethodName + ".xlsx", Q3_HEADER, rows)
        return validate_submission(3, path, base_q2=base or self.all_canceled, **kwargs)

    def assert_invalid(self, result, code=None):
        details = result["errors"][:4]
        self.assertFalse(result["passed"], details)
        self.assertTrue(result["errors"] or result["conflicts"], details)
        if code is not None:
            self.assertIn(code, {error["code"] for error in result["errors"]}, details)

    def test_known_q2_feasible_and_metrics(self):
        result = validate_submission(2, self.q2)
        self.assertTrue(result["passed"], result["errors"])
        self.assertEqual(result["conflicts"], [])
        self.assertEqual(list(result["summary"]["objective_tuple"]), [6, 0, 4, 126, 16, 34, 678, 86, 197])
        self.assertEqual(result["summary"]["occupied_area"], 15816)
        self.assertEqual(result["summary"]["demand_area"], 15816)

    def test_known_q3_feasible_and_area(self):
        result = validate_submission(3, self.q3, base_q2=self.q2)
        self.assertTrue(result["passed"], result["errors"])
        self.assertEqual(result["summary"]["added_count"], 136)
        self.assertEqual(result["summary"]["occupied_area"], 25608)
        self.assertEqual(result["summary"]["demand_area"], 25608)
        self.assertAlmostEqual(result["summary"]["utilization"], 25608 / 64300)

    def test_empty_q2_keeps_original_conflicts(self):
        path = self.table("empty_q2.xlsx", Q2_HEADER, [])
        result = validate_submission(2, path)
        self.assert_invalid(result)
        self.assertTrue(result["conflicts"])
        self.assertEqual(result["summary"]["occupied_area"], 14698)
        self.assertEqual(result["summary"]["demand_area"], 16680)

    def test_empty_q3_with_feasible_base_is_feasible(self):
        result = self.q3_rows([], base=self.q2)
        self.assertTrue(result["passed"], result["errors"])
        self.assertEqual(result["summary"]["added_count"], 0)
        self.assertEqual(result["summary"]["occupied_area"], 15816)

    def test_duplicate_q2_id_rejected(self):
        result = validate_submission(2, APP / "examples/result2_重复编号示例.xlsx")
        self.assert_invalid(result, "duplicate_id")

    def test_unknown_q2_id_rejected(self):
        path = self.table("unknown.xlsx", Q2_HEADER, self.all_canceled_rows + [["A999", None, None, "是"]])
        self.assert_invalid(validate_submission(2, path), "unknown_id")

    def test_negative_frequency_not_lost(self):
        path = self.isolated_q2(["A002", "[-1,9)", None, None])
        result = validate_submission(2, path)
        self.assert_invalid(result, "frequency_boundary")
        messages = " ".join(error["message"] for error in result["errors"])
        self.assertTrue(any(word in messages for word in ["边界", "越界", "范围", "负", "0"]), messages)

    def test_negative_start_not_lost(self):
        result = self.q3_rows([[1, "[0,3)", "[-2,0)"]])
        self.assert_invalid(result, "time_boundary")

    def test_decimal_endpoint_rejected(self):
        result = self.q3_rows([[1, "[0.5,3.5)", "[0,2)"]])
        self.assert_invalid(result)

    def test_two_adjustment_actions_rejected(self):
        path = self.isolated_q2(["A001", "[81,91)", "[36,41)", None])
        self.assert_invalid(validate_submission(2, path), "multiple_actions")

    def test_unchanged_frequency_and_changed_time_is_one_action(self):
        path = self.isolated_q2(["A001", "[80,90)", "[36,41)", "否"])
        result = validate_submission(2, path)
        self.assertTrue(result["passed"], result["errors"])
        self.assertEqual(result["summary"]["objective_tuple"][3], 1)
        self.assertEqual(result["summary"]["objective_tuple"][6], 1)

    def test_frequency_shift_exactly_ten_is_valid(self):
        path = self.isolated_q2(["A001", "[90,100)", None, None])
        result = validate_submission(2, path)
        self.assertTrue(result["passed"], result["errors"])
        self.assertEqual(result["summary"]["objective_tuple"][6], 10)

    def test_bandwidth_change_rejected(self):
        path = self.isolated_q2(["A001", "[80,91)", None, None])
        self.assert_invalid(validate_submission(2, path), "changed_width")

    def test_duration_change_rejected(self):
        path = self.isolated_q2(["A001", None, "[35,41)", None])
        self.assert_invalid(validate_submission(2, path), "changed_duration")

    def test_repeat_count_change_rejected(self):
        path = self.table("count_change.xlsx", Q3_HEADER + ["使用次数"], [[1, "[0,3)", "[0,2)", 11]])
        result = validate_submission(3, path, base_q2=self.all_canceled)
        self.assert_invalid(result, "changed_count")

    def test_gap_change_rejected(self):
        path = self.table("gap_change.xlsx", Q3_HEADER + ["时间间隔"], [[1, "[0,3)", "[0,2)", 9]])
        result = validate_submission(3, path, base_q2=self.all_canceled)
        self.assert_invalid(result, "changed_gap")

    def test_frequency_shift_limit(self):
        result = validate_submission(2, APP / "examples/result2_频移超限示例.xlsx")
        self.assert_invalid(result, "frequency_shift_limit")

    def test_time_shift_limit(self):
        path = self.isolated_q2(["A001", None, "[41,46)", None])
        self.assert_invalid(validate_submission(2, path), "time_shift_limit")

    def test_q2_last_period_must_fit_horizon(self):
        path = self.isolated_q2(["C083", None, "[532,534)", None])
        result = validate_submission(2, path)
        self.assert_invalid(result, "time_boundary")

    def test_omitted_original_plan_must_fit_selected_horizon(self):
        rows = [row for row in self.all_canceled_rows if row[0] != "A001"]
        path = self.table("omitted_original_horizon.xlsx", Q2_HEADER, rows)
        too_short = validate_submission(2, path, horizon=169)
        self.assert_invalid(too_short, "time_boundary")
        self.assertTrue(any(error["code"] == "time_boundary" and error["id"] == "A001"
                            for error in too_short["errors"]))
        exact = validate_submission(2, path, horizon=170)
        self.assertTrue(exact["passed"], exact["errors"])

    def test_q3_last_period_must_fit_horizon(self):
        self.assert_invalid(self.q3_rows([[1, "[0,3)", "[532,534)"]]), "time_boundary")

    def test_q3_end_exactly_at_horizon_is_valid(self):
        result = self.q3_rows([[1, "[97,100)", "[531,533)"]])
        self.assertTrue(result["passed"], result["errors"])
        self.assertEqual(result["summary"]["occupied_area"], 72)

    def test_formula_id_is_not_silently_skipped(self):
        path = self.table("formula_id.xlsx", Q2_HEADER,
                          self.all_canceled_rows + [['="A999"', None, None, "是"]])
        self.assert_invalid(validate_submission(2, path), "formula_cell")

    def test_formula_interval_rejected(self):
        path = self.isolated_q2(["A001", '="[81,91)"', None, None])
        self.assert_invalid(validate_submission(2, path), "formula_cell")

    def test_formula_q3_serial_rejected(self):
        self.assert_invalid(self.q3_rows([["=1", "[0,3)", "[0,2)"]]), "formula_cell")

    def test_duplicate_q3_serial_rejected(self):
        self.assert_invalid(self.q3_rows([[1, "[0,3)", "[0,2)"], [1, "[3,6)", "[0,2)"]]), "duplicate_id")

    def test_q3_requires_base(self):
        result = validate_submission(3, self.q3)
        self.assert_invalid(result, "missing_base")

    def test_q3_invalid_base_cannot_pass(self):
        empty = self.table("bad_base.xlsx", Q2_HEADER, [])
        result = self.q3_rows([], base=empty)
        self.assert_invalid(result, "base_not_feasible")

    def test_overlapping_new_plans_rejected(self):
        result = validate_submission(3, APP / "examples/result3_新增重叠示例.xlsx", base_q2=self.q2)
        self.assert_invalid(result, "resource_conflict")
        self.assertTrue(result["conflicts"])

    def test_full_cycles_checked_beyond_first_interval(self):
        # 两个首次区间互不相交，但第二台首次使用与第一台第二次使用重合。
        result = self.q3_rows([[1, "[0,3)", "[0,2)"], [2, "[0,3)", "[10,12)"]])
        self.assert_invalid(result)
        self.assertTrue(result["conflicts"])
        self.assertEqual(result["summary"]["demand_area"], 144)
        self.assertEqual(result["summary"]["occupied_area"], 78)

    def test_frequency_endpoint_touch_is_not_conflict(self):
        result = self.q3_rows([[1, "[0,3)", "[0,2)"], [2, "[3,6)", "[0,2)"]])
        self.assertTrue(result["passed"], result["errors"])
        self.assertEqual(result["summary"]["occupied_area"], 144)

    def test_time_endpoint_touch_is_not_conflict(self):
        result = self.q3_rows([[1, "[0,3)", "[0,2)"], [2, "[0,3)", "[2,4)"]])
        self.assertTrue(result["passed"], result["errors"])
        self.assertEqual(result["summary"]["occupied_area"], 144)

    def test_named_sheet_selection(self):
        path = self.directory / "multiple_sheets.xlsx"
        book = load_workbook(self.q2)
        book.create_sheet("说明", 0).append(["请选择提交结果工作表"])
        book.save(path)
        result = validate_submission(2, path, sheet="提交结果")
        self.assertTrue(result["passed"], result["errors"])


if __name__ == "__main__":
    unittest.main()
